#!/usr/bin/env python3
"""Export the skill's simple Markdown dialect as editable Word; requires python-docx."""
import argparse, re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def rich(paragraph, text):
    pattern=r'(\*\*[^*]+\*\*|\[[^\]]+\]\([^\s)]+\)|`[^`]+`)'
    for token in re.split(pattern,text):
        m=re.fullmatch(r'\[([^\]]+)\]\(([^\s)]+)\)',token)
        if m:
            link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),paragraph.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
            col=OxmlElement('w:color'); col.set(qn('w:val'),'126E86');prop.append(col)
            u=OxmlElement('w:u');u.set(qn('w:val'),'single');prop.append(u);run.append(prop)
            t=OxmlElement('w:t');t.text=m[1];run.append(t);link.append(run);paragraph._p.append(link)
        elif token.startswith('**') and token.endswith('**'):
            run=paragraph.add_run(token[2:-2]);run.bold=True
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        elif token.startswith('`') and token.endswith('`'):
            paragraph.add_run(token[1:-1]).font.name='Menlo'
        else: paragraph.add_run(token)

def export(markdown, out, page_before=()):
    source=Path(markdown).resolve(); lines=source.read_text(encoding='utf-8').splitlines()
    d=Document(); sec=d.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=sec.bottom_margin=Inches(.72)
    sec.left_margin=sec.right_margin=Inches(.83)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption','List Bullet']:
        style=d.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=Pt(18);normal.paragraph_format.space_after=Pt(8)
    normal.paragraph_format.widow_control=True
    title=d.styles['Title'];title.font.size=Pt(21);title.font.bold=True
    title.paragraph_format.space_after=Pt(17);title.paragraph_format.line_spacing=Pt(29)
    for name,size in [('Heading 1',15),('Heading 2',12)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.space_before=Pt(16);st.paragraph_format.space_after=Pt(9)
        st.paragraph_format.keep_with_next=True
    d.styles['Caption'].font.size=Pt(9)
    d.styles['Caption'].font.italic=False
    d.styles['Caption'].font.bold=False
    d.styles['Caption'].paragraph_format.line_spacing=Pt(13)
    d.styles['Caption'].paragraph_format.space_after=Pt(10)
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1;continue
        h=re.match(r'^(#{1,3}) (.+)$',line);image=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if h:
            p=d.add_paragraph(style='Title' if len(h[1])==1 else 'Heading '+str(len(h[1])-1))
            if h[2] in page_before: p.paragraph_format.page_break_before=True
            rich(p,h[2])
        elif image:
            path=(source.parent/image[2]).resolve()
            if source.parent not in path.parents: raise ValueError('Image must be inside article directory')
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.line_spacing=1.0
            from PIL import Image
            w,h=Image.open(path).size
            width=min(6.58,4.45*w/h)
            shape=p.add_run().add_picture(str(path),width=Inches(width))
            shape._inline.docPr.set('descr',image[1]);shape._inline.docPr.set('title',image[1])
            cap=d.add_paragraph(image[1],style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('|') and i+1<len(lines) and re.fullmatch(r'[\s|:\-]+',lines[i+1]):
            rows=[[x.strip() for x in line.strip('|').split('|')]];i+=2
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]);i+=1
            table=d.add_table(rows=0,cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            widths=([3.28,1.1,1.1,1.1] if len(rows[0])==4 else [6.58/len(rows[0])]*len(rows[0]))
            for col,width in zip(table.columns,widths): col.width=Inches(width)
            for ri,row in enumerate(rows):
                cells=table.add_row().cells
                trpr=table.rows[-1]._tr.get_or_add_trPr()
                cant=OxmlElement('w:cantSplit');trpr.append(cant)
                if ri==0: trpr.append(OxmlElement('w:tblHeader'))
                for ci,(cell,text) in enumerate(zip(cells,row)):
                    cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr()
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
                    tcpr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        m=OxmlElement('w:'+edge);m.set(qn('w:w'),'95');m.set(qn('w:type'),'dxa');margins.append(m)
                    tcpr.append(margins)
                    if ri==0:
                        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'DDEBF7');tcpr.append(shading)
                    p=cell.paragraphs[0];rich(p,text);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(15)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size=Pt(10)
                        if ri==0:run.bold=True
            d.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        else:
            p=d.add_paragraph(style='List Bullet' if line.startswith('- ') else None)
            rich(p,line[2:] if line.startswith('- ') else line)
        i+=1
    d.core_properties.title=next((l[2:] for l in lines if l.startswith('# ')),'论文推文')
    d.core_properties.author='';d.core_properties.last_modified_by=''
    d.save(out)

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('markdown');ap.add_argument('--out',required=True);ap.add_argument('--page-before',action='append',default=[])
    a=ap.parse_args();export(a.markdown,a.out,a.page_before);print(Path(a.out).resolve())
